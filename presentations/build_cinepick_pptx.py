"""Build the CinePick college project PowerPoint."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = Path(r"C:\Users\rv012\.cursor\projects\c-Users-rv012-CinePick\assets")
OUT = ROOT / "CinePick_Presentation.pptx"
CACHE = ROOT / "presentations" / "_cache"
CACHE.mkdir(parents=True, exist_ok=True)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG = RGBColor(10, 10, 14)
CARD = RGBColor(22, 22, 30)
CARD_ALT = RGBColor(28, 28, 38)
RED = RGBColor(229, 9, 20)
RED_SOFT = RGBColor(255, 50, 58)
GOLD = RGBColor(232, 197, 71)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 176, 186)
SOFT = RGBColor(128, 128, 140)

FONT = "Segoe UI"
FONT_B = "Segoe UI Semibold"


def emu(inches: float) -> int:
    return int(Inches(inches))


def set_fill(shape, color: RGBColor, alpha: float | None = None) -> None:
    sp_pr = shape._element.spPr
    for child in list(sp_pr):
        if child.tag == qn("a:solidFill"):
            sp_pr.remove(child)
    solid = etree.SubElement(sp_pr, qn("a:solidFill"))
    srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    if alpha is not None:
        a = etree.SubElement(srgb, qn("a:alpha"))
        a.set("val", str(int(max(0.0, min(1.0, alpha)) * 100000)))


def no_line(shape) -> None:
    shape.line.fill.background()


def roundness(shape, value: float = 0.12) -> None:
    try:
        shape.adjustments[0] = value
    except Exception:
        pass


def add_rect(slide, l, t, w, h, color, alpha=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    set_fill(s, color, alpha)
    no_line(s)
    return s


def add_round(slide, l, t, w, h, color, alpha=None, adj=0.12):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    set_fill(s, color, alpha)
    no_line(s)
    roundness(s, adj)
    return s


def set_run(run, text, size, color, bold=False, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def add_text(slide, l, t, w, h, text, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size, color, bold, font)
    return box


def add_multitext(slide, l, t, w, h, lines, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = line.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(line.get("space_after", 4))
        set_run(
            p.add_run(),
            line["text"],
            line.get("size", 16),
            line.get("color", WHITE),
            line.get("bold", False),
            line.get("font", FONT),
        )
    return box


def add_picture(slide, path, l, t, w, h):
    return slide.shapes.add_picture(str(path), l, t, w, h)


def darken_image(src: Path, dest: Path, brightness=0.42, contrast=1.08):
    im = Image.open(src).convert("RGB")
    im = ImageEnhance.Brightness(im).enhance(brightness)
    im = ImageEnhance.Contrast(im).enhance(contrast)
    im = im.filter(ImageFilter.GaussianBlur(radius=0.6))
    im.save(dest, quality=92)


def crop_cover(src: Path, dest: Path, w: int, h: int, brightness=1.0, radius=0):
    im = Image.open(src).convert("RGB")
    src_ratio = im.width / im.height
    dest_ratio = w / h
    if src_ratio > dest_ratio:
        new_w = int(im.height * dest_ratio)
        left = (im.width - new_w) // 2
        im = im.crop((left, 0, left + new_w, im.height))
    else:
        new_h = int(im.width / dest_ratio)
        top = (im.height - new_h) // 2
        im = im.crop((0, top, im.width, top + new_h))
    im = im.resize((w, h), Image.Resampling.LANCZOS)
    if brightness != 1.0:
        im = ImageEnhance.Brightness(im).enhance(brightness)
    if radius:
        im = im.convert("RGBA")
        mask = Image.new("L", (w, h), 0)
        ImageDraw.Draw(mask).rounded_rectangle((0, 0, w - 1, h - 1), radius=radius, fill=255)
        im.putalpha(mask)
        dest = dest.with_suffix(".png")
        im.save(dest)
        return dest
    im.save(dest, quality=94)
    return dest


def footer(slide, n, total=10):
    add_rect(slide, 0, emu(7.42), SLIDE_W, emu(0.08), RED)
    add_text(slide, emu(0.5), emu(7.12), emu(8), emu(0.28), "CinePick  ·  2nd Year B.Tech Project", 11, SOFT)
    add_text(slide, emu(11.4), emu(7.12), emu(1.4), emu(0.28), f"{n:02d} / {total:02d}", 11, SOFT, align=PP_ALIGN.RIGHT)


def header(slide, kicker, title, subtitle=None):
    add_rect(slide, 0, 0, SLIDE_W, emu(0.08), RED)
    add_text(slide, emu(0.55), emu(0.22), emu(12), emu(0.28), kicker.upper(), 12, RED_SOFT, True)
    add_text(slide, emu(0.55), emu(0.46), emu(12), emu(0.48), title, 30, WHITE, True, font=FONT_B)
    if subtitle:
        add_text(slide, emu(0.55), emu(0.96), emu(12), emu(0.32), subtitle, 14, MUTED)


def blank_dark(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, BG)
    return slide


def slide_title(prs):
    slide = blank_dark(prs)
    bg = CACHE / "title_bg.jpg"
    crop_cover(ASSETS / "cinepick_title_bg.jpg", bg, 1920, 1080, brightness=0.55)
    add_picture(slide, bg, 0, 0, SLIDE_W, SLIDE_H)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, RGBColor(0, 0, 0), 0.48)
    add_rect(slide, 0, 0, emu(0.18), SLIDE_H, RED)

    posters = crop_cover(ASSETS / "cinepick_posters.jpg", CACHE / "title_posters.png", 1400, 900, brightness=0.78, radius=48)
    add_picture(slide, posters, emu(7.35), emu(1.35), emu(5.5), emu(4.55))

    add_round(slide, emu(0.55), emu(1.55), emu(2.35), emu(0.38), RED, adj=0.5)
    add_text(slide, emu(0.55), emu(1.58), emu(2.35), emu(0.34), "COLLEGE PROJECT", 11, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

    add_text(slide, emu(0.55), emu(2.15), emu(7.0), emu(1.15), "CinePick", 66, WHITE, True, font=FONT_B)
    add_text(
        slide,
        emu(0.55),
        emu(3.25),
        emu(6.7),
        emu(1.1),
        "AI Powered Personalized\nMovie Recommendation System",
        24,
        RGBColor(230, 230, 235),
        False,
    )
    add_text(
        slide,
        emu(0.55),
        emu(4.55),
        emu(6.5),
        emu(0.7),
        "A smart movie assistant that understands your taste\nand helps you pick what to watch — faster.",
        15,
        MUTED,
    )

    add_round(slide, emu(0.55), emu(5.5), emu(2.2), emu(0.48), RED, adj=0.5)
    add_text(slide, emu(0.55), emu(5.52), emu(2.2), emu(0.44), "Your personal cinema", 13, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_round(slide, emu(2.9), emu(5.5), emu(2.35), emu(0.48), WHITE, 0.12, 0.5)
    add_text(slide, emu(2.9), emu(5.52), emu(2.35), emu(0.44), "2nd Year B.Tech", 13, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

    add_text(slide, emu(0.55), emu(6.85), emu(8), emu(0.3), "Department of Computer Science  ·  Academic Presentation", 12, SOFT)


def slide_problem(prs):
    slide = blank_dark(prs)
    header(slide, "01  ·  The Problem", "Finding a good movie should not feel like work")
    cards = [
        ("01", "Too much time", "Scrolling for 30 minutes often takes longer than watching the movie."),
        ("02", "Too many choices", "Thousands of titles create confusion, not excitement."),
        ("03", "Not personal", "Most platforms push trending films — not what you actually enjoy."),
    ]
    x = 0.5
    for num, title, body in cards:
        add_round(slide, emu(x), emu(1.7), emu(3.9), emu(4.55), CARD, adj=0.08)
        add_round(slide, emu(x + 0.28), emu(2.05), emu(0.72), emu(0.42), RED, adj=0.4)
        add_text(slide, emu(x + 0.28), emu(2.07), emu(0.72), emu(0.38), num, 14, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, emu(x + 0.28), emu(2.75), emu(3.35), emu(1.1), title, 26, WHITE, True, font=FONT_B)
        add_text(slide, emu(x + 0.28), emu(3.9), emu(3.35), emu(1.6), body, 16, MUTED)
        x += 4.15
    footer(slide, 2)


def slide_solution(prs):
    slide = blank_dark(prs)
    header(slide, "02  ·  Our Solution", "CinePick is your personal movie assistant")
    img = crop_cover(ASSETS / "cinepick_assistant.jpg", CACHE / "solution.png", 1200, 900, 0.95, radius=42)
    add_picture(slide, img, emu(7.55), emu(1.6), emu(5.25), emu(4.85))

    points = [
        ("Understands your taste", "It learns the genres, moods, and stories you enjoy."),
        ("Recommends for you", "Suggestions are based on personal interest, not only trends."),
        ("Feels like an assistant", "Less scrolling. Faster decisions. Better movie nights."),
    ]
    y = 1.65
    for title, body in points:
        add_round(slide, emu(0.5), emu(y), emu(6.7), emu(1.45), CARD, adj=0.1)
        add_round(slide, emu(0.72), emu(y + 0.52), emu(0.14), emu(0.42), RED, adj=0.5)
        add_text(slide, emu(1.1), emu(y + 0.22), emu(5.8), emu(0.45), title, 20, WHITE, True, font=FONT_B)
        add_text(slide, emu(1.1), emu(y + 0.7), emu(5.8), emu(0.55), body, 15, MUTED)
        y += 1.6
    footer(slide, 3)


def slide_flow(prs):
    slide = blank_dark(prs)
    header(slide, "03  ·  How You Use It", "From opening the app to a perfect pick")
    steps = [
        ("1", "Open CinePick", "Start your movie night"),
        ("2", "Quick survey", "Share genres, mood & taste"),
        ("3", "Search or like", "Tell the system what you enjoy"),
        ("4", "Taste is learned", "CinePick understands you"),
        ("5", "Get picks", "Personalized recommendations"),
    ]
    x = 0.4
    for i, (num, title, body) in enumerate(steps):
        add_round(slide, emu(x), emu(2.15), emu(2.2), emu(3.55), CARD, adj=0.1)
        add_round(slide, emu(x + 0.7), emu(2.45), emu(0.8), emu(0.8), RED, adj=0.5)
        add_text(slide, emu(x + 0.7), emu(2.5), emu(0.8), emu(0.7), num, 24, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, emu(x + 0.12), emu(3.5), emu(1.96), emu(0.9), title, 16, WHITE, True, PP_ALIGN.CENTER, font=FONT_B)
        add_text(slide, emu(x + 0.12), emu(4.4), emu(1.96), emu(0.85), body, 13, MUTED, False, PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, emu(x + 2.22), emu(3.7), emu(0.28), emu(0.22))
            set_fill(arrow, RED)
            no_line(arrow)
        x += 2.58
    add_text(slide, emu(0.55), emu(6.0), emu(12.2), emu(0.4), "Simple idea: you tell CinePick what you like  →  it finds movies that match.", 15, GOLD, False, PP_ALIGN.CENTER)
    footer(slide, 4)


def movie_card(slide, x, y, w, h, image, title, tag):
    add_round(slide, emu(x), emu(y), emu(w), emu(h), CARD, adj=0.08)
    if image:
        dest = crop_cover(image, CACHE / f"card_{title.replace(' ', '_')}.png", 520, 720, 0.95, radius=28)
        add_picture(slide, dest, emu(x + 0.1), emu(y + 0.1), emu(w - 0.2), emu(h - 0.92))
    add_text(slide, emu(x + 0.08), emu(y + h - 0.78), emu(w - 0.16), emu(0.34), title, 12, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, emu(x + 0.08), emu(y + h - 0.46), emu(w - 0.16), emu(0.28), tag, 10, MUTED, False, PP_ALIGN.CENTER)


def slide_example(prs):
    slide = blank_dark(prs)
    header(slide, "04  ·  Example", "How CinePick turns likes into suggestions")

    add_text(slide, emu(0.42), emu(1.48), emu(4.6), emu(0.32), "USER LIKES", 12, RED_SOFT, True)
    likes = [
        (ASSETS / "poster_space.jpg", "Interstellar", "Space epic"),
        (ASSETS / "poster_city.jpg", "Inception", "Mind-bending"),
        (ASSETS / "poster_magic.jpg", "Dark", "Dark thriller"),
    ]
    x = 0.42
    for img, title, tag in likes:
        movie_card(slide, x, 1.85, 1.55, 4.45, img, title, tag)
        x += 1.68

    add_round(slide, emu(5.52), emu(2.15), emu(2.22), emu(3.85), CARD, adj=0.1)
    add_text(slide, emu(5.57), emu(2.32), emu(2.12), emu(0.36), "SYSTEM READS", 11, RED_SOFT, True, PP_ALIGN.CENTER)
    insights = [
        ("Sci-Fi", "Loves big ideas"),
        ("Complex stories", "Enjoys twists"),
        ("Thriller", "Wants tension"),
    ]
    y = 2.82
    for title, body in insights:
        add_round(slide, emu(5.7), emu(y), emu(1.86), emu(0.88), RGBColor(40, 18, 22), adj=0.2)
        add_text(slide, emu(5.74), emu(y + 0.1), emu(1.78), emu(0.34), title, 13, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, emu(5.74), emu(y + 0.44), emu(1.78), emu(0.32), body, 11, MUTED, False, PP_ALIGN.CENTER)
        y += 0.98

    add_text(slide, emu(7.95), emu(1.48), emu(4.9), emu(0.32), "THEN RECOMMENDS", 12, GOLD, True)
    recs = [
        (ASSETS / "poster_arrival.jpg", "Arrival", "Thoughtful Sci-Fi"),
        (ASSETS / "poster_tenet.jpg", "Tenet", "Puzzle thriller"),
        (ASSETS / "poster_prestige.jpg", "The Prestige", "Smart mystery"),
    ]
    x = 7.95
    for img, title, tag in recs:
        movie_card(slide, x, 1.85, 1.55, 4.45, img, title, tag)
        x += 1.68
    footer(slide, 5)


def slide_features(prs):
    slide = blank_dark(prs)
    header(slide, "05  ·  Features", "Everything you need for a better movie night")
    items = [
        ("★", "Personalized picks", "Movies matched to your taste, not the crowd."),
        ("⌕", "Movie search", "Find any film quickly and explore similar titles."),
        ("◎", "Similar movies", "Liked one film? Discover others in the same vibe."),
        ("☰", "Smart filters", "Filter by genre, mood, rating, and language."),
        ("♡", "Watchlist", "Save movies for later and build your list."),
        ("✦", "Ratings", "Rate what you watch so suggestions keep improving."),
    ]
    positions = [(0.5, 1.65), (4.7, 1.65), (8.9, 1.65), (0.5, 4.25), (4.7, 4.25), (8.9, 4.25)]
    for (x, y), (icon, title, body) in zip(positions, items):
        add_round(slide, emu(x), emu(y), emu(3.9), emu(2.3), CARD, adj=0.1)
        add_round(slide, emu(x + 0.25), emu(y + 0.28), emu(0.55), emu(0.55), RED, adj=0.25)
        add_text(slide, emu(x + 0.25), emu(y + 0.3), emu(0.55), emu(0.5), icon, 16, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, emu(x + 0.95), emu(y + 0.32), emu(2.7), emu(0.5), title, 16, WHITE, True, font=FONT_B)
        add_text(slide, emu(x + 0.28), emu(y + 1.1), emu(3.35), emu(0.9), body, 14, MUTED)
    footer(slide, 6)


def slide_internal(prs):
    slide = blank_dark(prs)
    header(slide, "06  ·  Inside the System", "A simple path from your input to a movie pick")
    stages = [
        ("User Input", "Survey, search\nand likes"),
        ("Backend", "Processes\nrequests"),
        ("Database", "Stores users\nand movies"),
        ("AI Engine", "Finds matching\ntaste patterns"),
        ("Suggestions", "Personalized\nmovie list"),
    ]
    x = 0.4
    for i, (title, body) in enumerate(stages):
        add_round(slide, emu(x), emu(2.35), emu(2.15), emu(3.0), CARD if i != 3 else RGBColor(42, 12, 16), adj=0.1)
        add_round(slide, emu(x + 0.75), emu(2.6), emu(0.65), emu(0.42), RED, adj=0.4)
        add_text(slide, emu(x + 0.75), emu(2.62), emu(0.65), emu(0.38), str(i + 1), 14, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, emu(x + 0.1), emu(3.25), emu(1.95), emu(0.7), title, 16, WHITE, True, PP_ALIGN.CENTER, font=FONT_B)
        add_text(slide, emu(x + 0.12), emu(4.05), emu(1.9), emu(0.9), body, 13, MUTED, False, PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            chev = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, emu(x + 2.18), emu(3.7), emu(0.32), emu(0.24))
            set_fill(chev, GOLD if i == 3 else RED)
            no_line(chev)
        x += 2.58
    add_round(slide, emu(1.4), emu(5.7), emu(10.5), emu(0.7), CARD_ALT, adj=0.4)
    add_text(
        slide,
        emu(1.4),
        emu(5.78),
        emu(10.5),
        emu(0.55),
        "You  →  CinePick server  →  stored profile  →  recommendation logic  →  movies for you",
        14,
        WHITE,
        False,
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
    )
    footer(slide, 7)


def slide_tech(prs):
    slide = blank_dark(prs)
    header(slide, "07  ·  Technology", "Built with modern, practical tools")
    items = [
        ("React", "User interface", "Creates the screens people see and click."),
        ("Tailwind CSS", "Modern design", "Keeps the app clean, dark, and Netflix-like."),
        ("Node.js", "Backend", "Handles login, requests, and recommendations."),
        ("MongoDB", "Database", "Stores user accounts, likes, and watchlists."),
        ("TMDB API", "Movie details", "Fetches posters, ratings, and film information."),
        ("AI Algorithm", "Personal picks", "Matches movies to each user’s unique taste."),
    ]
    positions = [(0.5, 1.65), (4.7, 1.65), (8.9, 1.65), (0.5, 4.25), (4.7, 4.25), (8.9, 4.25)]
    for (x, y), (name, role, body) in zip(positions, items):
        add_round(slide, emu(x), emu(y), emu(3.9), emu(2.3), CARD, adj=0.1)
        add_text(slide, emu(x + 0.28), emu(y + 0.25), emu(3.35), emu(0.4), name, 20, WHITE, True, font=FONT_B)
        add_text(slide, emu(x + 0.28), emu(y + 0.7), emu(3.35), emu(0.32), role.upper(), 11, RED_SOFT, True)
        add_text(slide, emu(x + 0.28), emu(y + 1.15), emu(3.35), emu(0.85), body, 14, MUTED)
    footer(slide, 8)


def slide_future(prs):
    slide = blank_dark(prs)
    header(slide, "08  ·  Future Scope", "Where CinePick can go next")
    img = crop_cover(ASSETS / "cinepick_ai.jpg", CACHE / "future.png", 1100, 900, 0.9, radius=42)
    add_picture(slide, img, emu(8.55), emu(1.7), emu(4.25), emu(4.7))
    items = [
        ("01", "Smarter AI model", "A stronger machine learning model for even better matching."),
        ("02", "Movie chatbot", "Ask in plain English: “something like Interstellar, but shorter.”"),
        ("03", "Voice search", "Say a mood or movie name and get instant suggestions."),
        ("04", "Deeper personalization", "Learn from ratings, watch time, and changing taste."),
    ]
    y = 1.6
    for num, title, body in items:
        add_round(slide, emu(0.5), emu(y), emu(7.75), emu(1.2), CARD, adj=0.12)
        add_round(slide, emu(0.72), emu(y + 0.35), emu(0.7), emu(0.5), RED, adj=0.35)
        add_text(slide, emu(0.72), emu(y + 0.38), emu(0.7), emu(0.44), num, 13, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, emu(1.6), emu(y + 0.18), emu(6.4), emu(0.4), title, 18, WHITE, True, font=FONT_B)
        add_text(slide, emu(1.6), emu(y + 0.62), emu(6.4), emu(0.42), body, 14, MUTED)
        y += 1.3
    footer(slide, 9)


def slide_conclusion(prs):
    slide = blank_dark(prs)
    bg = CACHE / "end_bg.jpg"
    crop_cover(ASSETS / "cinepick_title_bg.jpg", bg, 1920, 1080, 0.38)
    add_picture(slide, bg, 0, 0, SLIDE_W, SLIDE_H)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, RGBColor(0, 0, 0), 0.55)
    add_rect(slide, 0, 0, emu(0.18), SLIDE_H, RED)

    add_text(slide, emu(0.7), emu(1.55), emu(12), emu(0.35), "CONCLUSION", 13, RED_SOFT, True, PP_ALIGN.CENTER)
    add_text(slide, emu(1.2), emu(2.05), emu(10.9), emu(1.5), "CinePick helps you discover movies easily\nby understanding your personal taste.", 28, WHITE, True, PP_ALIGN.CENTER, font=FONT_B)
    add_text(
        slide,
        emu(2.2),
        emu(3.7),
        emu(8.9),
        emu(0.8),
        "Less scrolling. More watching. A movie night that actually feels like yours.",
        16,
        MUTED,
        False,
        PP_ALIGN.CENTER,
    )

    pills = ["Personalized", "Simple to use", "Ready to grow"]
    x = 3.35
    for p in pills:
        add_round(slide, emu(x), emu(4.75), emu(2.15), emu(0.5), RED, 0.95, 0.5)
        add_text(slide, emu(x), emu(4.78), emu(2.15), emu(0.44), p, 13, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        x += 2.35

    add_text(slide, emu(0.7), emu(5.7), emu(12), emu(0.45), "Thank you", 32, WHITE, True, PP_ALIGN.CENTER, font=FONT_B)
    add_text(slide, emu(0.7), emu(6.2), emu(12), emu(0.35), "Questions are welcome.", 16, GOLD, False, PP_ALIGN.CENTER)
    footer(slide, 10)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide_title(prs)
    slide_problem(prs)
    slide_solution(prs)
    slide_flow(prs)
    slide_example(prs)
    slide_features(prs)
    slide_internal(prs)
    slide_tech(prs)
    slide_future(prs)
    slide_conclusion(prs)
    prs.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
